"""
Ejercicio 3 - Rachas

Arma la base SQLite desde rachas.xlsx y corre la consulta de rachas.

Uso:
    python rachas.py cargar
    python rachas.py consultar --fecha-base 2024-12-31 --n 3
    python rachas.py consultar --fecha-base 2024-06-15 --n 4 --salida resultados/rachas.csv
    python rachas.py detalle --identificacion R7Q4Z9AULIORDJQ9D --fecha-base 2024-12-31
    python rachas.py calidad

Python aca solo hace de pegamento: lee el Excel, lo mete en staging y corre
los .sql en orden. Las reglas de negocio estan todas en SQL (es lo que pide el
ejercicio y asi se puede llevar a otro motor sin reescribir nada).
"""

import argparse
import csv
import sqlite3
import sys
from contextlib import closing
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
SQL_DIR = BASE_DIR / "sql"
DEFAULT_EXCEL = BASE_DIR / "data" / "rachas.xlsx"
DEFAULT_DB = BASE_DIR / "rachas.db"

# hoja del Excel -> (tabla de staging, columnas que se esperan)
SHEETS = {
    "historia": ("stg_historia", ("identificacion", "corte_mes", "saldo")),
    "retiros": ("stg_retiros", ("identificacion", "fecha_retiro")),
}


# --- Excel -------------------------------------------------------------------
def _to_staging_value(value):
    """Fechas a texto YYYY-MM-DD (lo que entiende SQLite). Lo demas pasa igual,
    la validacion se hace en SQL."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def read_sheet(workbook, sheet_name, expected_columns):
    """Devuelve las filas de la hoja como (fila_origen, col1, col2, ...).

    Busco las columnas por nombre y no por posicion, por si alguien las
    reordena en el Excel.
    """
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"El archivo no tiene la hoja '{sheet_name}'")

    rows = workbook[sheet_name].iter_rows(values_only=True)
    header = [str(col).strip().lower() if col is not None else "" for col in next(rows)]

    missing = [col for col in expected_columns if col not in header]
    if missing:
        raise ValueError(f"La hoja '{sheet_name}' no tiene las columnas: {missing}")
    positions = [header.index(col) for col in expected_columns]

    records = []
    # start=2 porque la fila 1 es el encabezado
    for excel_row, row in enumerate(rows, start=2):
        if row is None or all(cell is None for cell in row):
            continue  # filas vacias al final
        values = [_to_staging_value(row[pos]) if pos < len(row) else None for pos in positions]
        records.append((excel_row, *values))
    return records


# --- Base de datos -------------------------------------------------------------
def connect(db_path):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def run_sql_file(conn, file_name):
    conn.executescript((SQL_DIR / file_name).read_text(encoding="utf-8"))


def load_staging(conn, excel_path):
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        for sheet_name, (table, columns) in SHEETS.items():
            records = read_sheet(workbook, sheet_name, columns)
            placeholders = ", ".join("?" for _ in range(len(columns) + 1))
            conn.executemany(
                f"INSERT INTO {table} (fila_origen, {', '.join(columns)}) VALUES ({placeholders})",
                records,
            )
    finally:
        workbook.close()


def build_database(excel_path, db_path):
    """Crea la base desde cero: tablas -> staging -> calidad -> limpia."""
    # ojo: "with sqlite3.connect()" hace commit pero NO cierra la conexion.
    # En Windows eso me dejaba el .db bloqueado, de ahi el closing()
    with closing(connect(db_path)) as conn, conn:
        run_sql_file(conn, "01_ddl.sql")
        load_staging(conn, excel_path)
        run_sql_file(conn, "02_calidad_datos.sql")
        run_sql_file(conn, "03_limpieza.sql")
        return load_summary(conn)


def load_summary(conn):
    def count(table):
        return conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]

    incidents = conn.execute(
        """
        SELECT regla, severidad, accion, COUNT(*) AS cantidad
        FROM dq_incidencias
        GROUP BY regla, severidad, accion
        ORDER BY regla
        """
    ).fetchall()
    return {
        "stg_historia": count("stg_historia"),
        "historia": count("historia"),
        "stg_retiros": count("stg_retiros"),
        "retiros": count("retiros"),
        "incidencias": [dict(row) for row in incidents],
    }


def query_rachas(conn, fecha_base, n):
    sql = (SQL_DIR / "04_rachas.sql").read_text(encoding="utf-8")
    return conn.execute(sql, {"fecha_base": fecha_base, "n": n}).fetchall()


def query_detalle(conn, identificacion, fecha_base):
    sql = (SQL_DIR / "05_detalle_cliente.sql").read_text(encoding="utf-8")
    return conn.execute(sql, {"identificacion": identificacion, "fecha_base": fecha_base}).fetchall()


def query_incidencias(conn):
    return conn.execute(
        """
        SELECT tabla_origen, fila_origen, identificacion, regla, severidad, accion, detalle
        FROM dq_incidencias
        ORDER BY id_incidencia
        """
    ).fetchall()


# --- Salida ------------------------------------------------------------------
def print_table(rows):
    if not rows:
        print("(sin resultados)")
        return
    columns = rows[0].keys()
    widths = [max(len(str(col)), *(len(str(row[col])) for row in rows)) for col in columns]
    print("  ".join(str(col).ljust(w) for col, w in zip(columns, widths)))
    print("  ".join("-" * w for w in widths))
    for row in rows:
        print("  ".join(str(row[col]).ljust(w) for col, w in zip(columns, widths)))


def write_csv(rows, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        if rows:
            writer.writerow(rows[0].keys())
            writer.writerows(tuple(row) for row in rows)


# --- CLI ---------------------------------------------------------------------
def iso_date(value):
    """Para argparse: que la fecha venga como YYYY-MM-DD."""
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise argparse.ArgumentTypeError(f"'{value}' no es una fecha YYYY-MM-DD") from exc


def positive_int(value):
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("n debe ser un entero mayor o igual a 1")
    return number


def build_parser():
    parser = argparse.ArgumentParser(description="Ejercicio 3 - Rachas de nivel de deuda")
    parser.add_argument("--db", default=DEFAULT_DB, type=Path, help="Ruta de la base SQLite")
    sub = parser.add_subparsers(dest="comando", required=True)

    cargar = sub.add_parser("cargar", help="Crea la base de datos a partir del Excel")
    cargar.add_argument("--excel", default=DEFAULT_EXCEL, type=Path)

    consultar = sub.add_parser("consultar", help="Ejecuta la consulta de rachas")
    consultar.add_argument("--fecha-base", required=True, type=iso_date)
    consultar.add_argument("--n", required=True, type=positive_int)
    consultar.add_argument("--salida", type=Path, help="CSV donde guardar el resultado")

    detalle = sub.add_parser("detalle", help="Serie mensual de un cliente")
    detalle.add_argument("--identificacion", required=True)
    detalle.add_argument("--fecha-base", required=True, type=iso_date)

    sub.add_parser("calidad", help="Lista las incidencias de calidad de datos")
    return parser


def main(argv=None):
    args = build_parser().parse_args(argv)

    if args.comando == "cargar":
        summary = build_database(args.excel, args.db)
        print(f"Base creada en {args.db}")
        print(f"  historia: {summary['stg_historia']} filas leídas -> {summary['historia']} cargadas")
        print(f"  retiros : {summary['stg_retiros']} filas leídas -> {summary['retiros']} cargadas")
        print("  incidencias de calidad:")
        for item in summary["incidencias"]:
            print(f"    - {item['regla']} [{item['severidad']}/{item['accion']}]: {item['cantidad']}")
        return 0

    if not args.db.exists():
        print(f"No existe la base {args.db}. Corre primero: python rachas.py cargar", file=sys.stderr)
        return 1

    with closing(connect(args.db)) as conn:
        if args.comando == "consultar":
            rows = query_rachas(conn, args.fecha_base, args.n)
            print_table(rows)
            print(f"\n{len(rows)} clientes con racha >= {args.n} a la fecha base {args.fecha_base}")
            if args.salida:
                write_csv(rows, args.salida)
                print(f"Resultado guardado en {args.salida}")
        elif args.comando == "detalle":
            print_table(query_detalle(conn, args.identificacion, args.fecha_base))
        elif args.comando == "calidad":
            print_table(query_incidencias(conn))
    return 0


if __name__ == "__main__":
    sys.exit(main())
